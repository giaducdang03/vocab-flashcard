import csv
import io
from typing import Any

from openpyxl import load_workbook


def parse_import_file(file_bytes: bytes, filename: str) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    lower_name = filename.lower()

    if lower_name.endswith(".csv"):
        text = file_bytes.decode("utf-8-sig")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            rows.append(normalize_row(row))
        return rows

    if lower_name.endswith(".xlsx"):
        workbook = load_workbook(filename=io.BytesIO(file_bytes), read_only=True)
        sheet = workbook.active
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            data = dict(zip(headers, row, strict=False))
            rows.append(normalize_row(data))
        return rows

    raise ValueError("Unsupported file type")


def normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    clean = {str(k).strip(): (v if v is not None else "") for k, v in row.items()}
    front = str(clean.get("front_text", "")).strip()
    phonetic = str(clean.get("phonetic", "") or "").strip() or None
    back = str(clean.get("back_text", "")).strip()
    example = str(clean.get("example", "") or "").strip() or None

    synonyms_value = str(clean.get("synonyms", "") or "").strip()
    synonyms = []
    if synonyms_value:
        for item in synonyms_value.split(";"):
            text = item.strip()
            if not text:
                continue
            if "/" in text:
                word, phon = text.split("/", 1)
                synonyms.append({"word": word.strip(), "phonetic": phon.strip(" /") or None})
            else:
                synonyms.append({"word": text, "phonetic": None})

    card_type = str(clean.get("card_type", "") or "").strip().lower()
    if not card_type:
        card_type = "vocab" if synonyms else "collocation"

    return {
        "card_type": card_type,
        "front_text": front,
        "front_phonetic": phonetic,
        "back_text": back,
        "example": example,
        "synonyms": synonyms,
    }
